import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_shading(cell, color_hex):
    """Aplica cor de fundo em célula de tabela."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Ajusta padding da célula."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def markdown_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()
    
    # Configuração de Margens da Página (2 cm)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Estilos de cabeçalho
    in_table = False
    table_lines = []

    def flush_table():
        nonlocal in_table, table_lines
        if not table_lines:
            return
        
        # Filtra linhas de separador de tabela (|---|---|)
        rows_data = []
        for line in table_lines:
            stripped = line.strip()
            if re.match(r'^\|[\s\-:|]+\|$', stripped):
                continue
            cells = [c.strip() for c in stripped.strip('|').split('|')]
            rows_data.append(cells)
        
        if rows_data:
            num_cols = max(len(r) for r in rows_data)
            table = doc.add_table(rows=len(rows_data), cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True
            
            for row_idx, row in enumerate(rows_data):
                for col_idx in range(num_cols):
                    cell = table.cell(row_idx, col_idx)
                    text = row[col_idx] if col_idx < len(row) else ""
                    # Remove formatações básicas de markdown no texto da célula
                    clean_text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
                    clean_text = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', clean_text)
                    cell.text = clean_text
                    
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.15
                    
                    if row_idx == 0:
                        set_cell_shading(cell, "0B5394") # Azul corporativo
                        for run in p.runs:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
                            run.font.name = "Arial"
                            run.font.size = Pt(9.5)
                    else:
                        bg_color = "F3F3F3" if row_idx % 2 == 1 else "FFFFFF"
                        set_cell_shading(cell, bg_color)
                        for run in p.runs:
                            run.font.name = "Arial"
                            run.font.size = Pt(9)
                            run.font.color.rgb = RGBColor(40, 40, 40)
            doc.add_paragraph() # Espaço após a tabela
        
        table_lines = []
        in_table = False

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Detecção de Tabelas
        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_lines.append(stripped)
            i += 1
            continue
        elif in_table:
            flush_table()

        if not stripped:
            i += 1
            continue

        # Título 1 (# Título)
        if stripped.startswith('# '):
            text = stripped[2:].strip()
            # Remove emojis se quiser ou mantém
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(text)
            run.font.name = 'Arial'
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(11, 83, 148) # Azul escuro
            i += 1
            continue

        # Título 2 (## Título)
        if stripped.startswith('## '):
            text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(text)
            run.font.name = 'Arial'
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 70, 110)
            i += 1
            continue

        # Título 3 (### Título)
        if stripped.startswith('### '):
            text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(text)
            run.font.name = 'Arial'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(60, 60, 60)
            i += 1
            continue

        # Citação / Callout (> texto)
        if stripped.startswith('> '):
            text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(text)
            run.font.name = 'Arial'
            run.font.size = Pt(9.5)
            run.font.italic = True
            run.font.color.rgb = RGBColor(80, 80, 80)
            i += 1
            continue

        # Divisor (---)
        if stripped == '---':
            i += 1
            continue

        # Lista com marcadores (* ou -)
        if stripped.startswith('* ') or stripped.startswith('- '):
            text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.15
            
            # Trata negrito dentro do marcador
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    run = p.add_run(part)
                run.font.name = 'Arial'
                run.font.size = Pt(10)
            i += 1
            continue

        # Lista numerada (1., 2., etc)
        num_match = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if num_match:
            num = num_match.group(1)
            text = num_match.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    run = p.add_run(part)
                run.font.name = 'Arial'
                run.font.size = Pt(10)
            i += 1
            continue

        # Bloco de código (``` ... ```)
        if stripped.startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i].rstrip())
                i += 1
            i += 1 # Pula o fechamento ```
            
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run('\n'.join(code_lines))
            run.font.name = 'Consolas'
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(30, 30, 30)
            continue

        # Parágrafo comum
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        
        parts = re.split(r'(\*\*.*?\*\*)', stripped)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
            else:
                # Remove links markdown [texto](url) mantendo o texto
                clean_part = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', part)
                run = p.add_run(clean_part)
            run.font.name = 'Arial'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(30, 30, 30)
            
        i += 1

    if in_table:
        flush_table()

    doc.save(docx_path)
    print(f"[OK] Gerado: {os.path.basename(docx_path)}")

def main():
    src_dir = r"c:\Users\SAAV166\Documents\Agente\base_conhecimento\ambiente_saavedra"
    dest_dir = r"G:\Drives compartilhados\Informatica\Documentacoes Sankhya - Melhorias e Consertos"
    
    os.makedirs(dest_dir, exist_ok=True)
    
    for f in sorted(os.listdir(src_dir)):
        if f.endswith('.md'):
            md_path = os.path.join(src_dir, f)
            docx_name = f.replace('.md', '.docx')
            docx_path = os.path.join(dest_dir, docx_name)
            try:
                markdown_to_docx(md_path, docx_path)
            except Exception as e:
                print(f"[ERRO] {f}: {e}")

if __name__ == '__main__':
    main()
